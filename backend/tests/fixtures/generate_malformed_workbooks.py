"""Run once locally to (re)generate malformed test fixtures used by
tests/integration/test_imports_api.py. Not part of the pytest suite itself."""
import openpyxl
from pathlib import Path

FIXTURES_DIR = Path(__file__).parent

# No worksheets at all is not achievable via openpyxl (a workbook always has
# at least one sheet), so "no useful sheet" is represented by a workbook
# whose only sheet has zero rows -- functionally identical to the empty-sheet
# case below and already covered by _validate_workbook_readable's has_data check.

# Empty sheet: a named sheet, but zero rows of any kind.
wb = openpyxl.Workbook()
wb.active.title = "EmptySheet"
wb.save(FIXTURES_DIR / "empty_sheet_workbook.xlsx")

# Headers only: header row present, but no data rows beneath it.
wb = openpyxl.Workbook()
ws = wb.active
ws.append(["Roll No", "Name", "Email", "Department", "Batch", "Group"])
wb.save(FIXTURES_DIR / "headers_only_workbook.xlsx")

# Unexpected structure: has data, but none of the columns match any known
# canonical field -- passes upload, but validate() should map nothing.
wb = openpyxl.Workbook()
ws = wb.active
ws.append(["Random Col A", "Random Col B", "Notes"])
ws.append(["x", "y", "z"])
wb.save(FIXTURES_DIR / "unexpected_structure_workbook.xlsx")

print("Fixtures generated in", FIXTURES_DIR)
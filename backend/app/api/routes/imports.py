from pathlib import Path
from tempfile import NamedTemporaryFile

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from sqlalchemy.orm import Session
import openpyxl

from app.core.database import get_db
from app.core.dependencies import require_role
from app.models.user import User
from app.schemas.import_schema import (
    ImportCreateResponse,
    ImportDetailResponse,
    ValidateImportResponse,
    CommitImportResponse,
)
from app.services.import_service import ImportService


router = APIRouter(
    prefix="/imports",
    tags=["Imports"],
)

faculty_required = require_role("faculty")

ALLOWED_EXTENSIONS = {".xlsx", ".xls"}
ALLOWED_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
    "application/octet-stream",
}
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024


def _validate_workbook_readable(file_path: str) -> None:
    try:
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file is not a valid or readable Excel workbook",
        ) from exc

    if not workbook.sheetnames:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded workbook contains no worksheets",
        )

    sheet = workbook[workbook.sheetnames[0]]
    has_data = False
    for row in sheet.iter_rows(max_row=2):
        if any(cell.value not in (None, "") for cell in row):
            has_data = True
            break

    workbook.close()

    if not has_data:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded workbook appears to be empty (no header or data rows found)",
        )


@router.post(
    "",
    response_model=ImportCreateResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_import(
    file: UploadFile = File(...),
    import_type: str = "student",
    db: Session = Depends(get_db),
    current_user: User = Depends(faculty_required),
):
    """
    Registers a new import: validates the upload boundary (extension,
    content type, size, workbook readability), saves the file, computes
    its hash, and creates an ImportJob record. Does NOT process rows yet
    — call POST /imports/{id}/validate next.
    """
    if not file.filename:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Filename is required")

    extension = Path(file.filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only Excel files (.xlsx, .xls) are supported",
        )

    if file.content_type and file.content_type not in ALLOWED_CONTENT_TYPES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported content type: {file.content_type}",
        )

    temporary_path = None
    try:
        total_size = 0
        with NamedTemporaryFile(suffix=extension, delete=False) as temporary_file:
            temporary_path = temporary_file.name
            while chunk := file.file.read(1024 * 1024):
                total_size += len(chunk)
                if total_size > MAX_FILE_SIZE_BYTES:
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=f"File exceeds maximum allowed size of {MAX_FILE_SIZE_BYTES // (1024 * 1024)} MB",
                    )
                temporary_file.write(chunk)

        if total_size == 0:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="The uploaded file is empty")

        _validate_workbook_readable(temporary_path)

        if import_type not in ("student", "department"):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported import_type: {import_type}",
            )

        import_job = ImportService.create_import(
            db=db,
            file_path=temporary_path,
            filename=file.filename,
            created_by=current_user.id,
            import_type=import_type,
        )

        return ImportCreateResponse(
            import_id=import_job.id,
            filename=import_job.filename,
            status=import_job.status,
            file_hash=import_job.file_hash,
        )

    except HTTPException:
        if temporary_path and Path(temporary_path).exists():
            Path(temporary_path).unlink(missing_ok=True)
        raise
    except Exception as exc:
        if temporary_path and Path(temporary_path).exists():
            Path(temporary_path).unlink(missing_ok=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Unable to process uploaded Excel file",
        ) from exc
    finally:
        file.file.close()

@router.get(
    "",
    response_model=list[ImportDetailResponse],
)
def list_imports(
    db: Session = Depends(get_db),
    current_user: User = Depends(faculty_required),
):
    import_jobs = ImportService.list_imports(db)
    return [
        ImportDetailResponse(
            import_id=job.id,
            filename=job.filename,
            status=job.status,
            created_by=job.created_by,
            created_at=job.created_at,
            completed_at=job.completed_at,
            total_rows=job.total_rows,
            valid_rows=job.valid_rows,
            invalid_rows=job.invalid_rows,
            duplicate_rows=job.duplicate_rows,
            error_message=job.error_message,
        )
        for job in import_jobs
    ]


@router.get(
    "/{import_id}",
    response_model=ImportDetailResponse,
)
def get_import(
    import_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(faculty_required),
):
    import_job = ImportService.get_import(db, import_id)
    if import_job is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Import not found")

    return ImportDetailResponse(
        import_id=import_job.id,
        filename=import_job.filename,
        status=import_job.status,
        created_by=import_job.created_by,
        created_at=import_job.created_at,
        completed_at=import_job.completed_at,
        total_rows=import_job.total_rows,
        valid_rows=import_job.valid_rows,
        invalid_rows=import_job.invalid_rows,
        duplicate_rows=import_job.duplicate_rows,
        error_message=import_job.error_message,
    )


@router.get("/{import_id}/errors")
def get_import_errors(
    import_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(faculty_required),
):
    import_job = ImportService.get_import(db, import_id)
    if import_job is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Import not found")

    if not import_job.validation_result:
        return {"import_id": import_job.id, "errors": []}

    payload = ImportService.get_preview(import_job)
    errors = [
        {
            "row": r["row"],
            "category": r["category"],
            "field_errors": r.get("field_errors", []),
            "reference_errors": r.get("reference_errors", []),
        }
        for r in payload["records"]
        if r["category"] != "VALID"
    ]
    return {"import_id": import_job.id, "errors": errors}


@router.post(
    "/{import_id}/validate",
    response_model=ValidateImportResponse,
)
def validate_import(
    import_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(faculty_required),
):
    """
    Runs the Excel pipeline and entity resolution, generates a preview.
    Writes ZERO student records — this only reads and computes.
    """
    import_job = ImportService.get_import(db, import_id)
    if import_job is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Import not found")

    try:
        result = ImportService.validate_import_dispatch(db, import_job)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Validation failed",
        ) from exc

    return ValidateImportResponse(
        import_id=import_job.id,
        status=import_job.status,
        mapping=result["mapping"],
        unmapped_columns=result["unmapped_columns"],
        ambiguous_columns=result.get("ambiguous_columns", []),
        column_status=result.get("column_status", []),
        summary=result["summary"],
        records=result["records"],
    )


@router.post(
    "/{import_id}/commit",
    response_model=CommitImportResponse,
)
def commit_import(
    import_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(faculty_required),
):
    """
    Commits a validated import: inserts only VALID records as students,
    inside a transaction. Rolls back entirely on failure.
    """
    import_job = ImportService.get_import(db, import_id)
    if import_job is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Import not found")

    try:
        result = ImportService.commit_import_dispatch(db, import_job)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Commit failed",
        ) from exc

    return CommitImportResponse(**result)
